#!/usr/bin/env python3
"""
Solution for finding column headers of smallest value greater than zero in each row.
"""

import openpyxl
import argparse
from pathlib import Path


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', required=True, help='Input Excel file')
    parser.add_argument('--output', required=True, help='Output Excel file')
    args = parser.parse_args()

    # Load workbook
    wb = openpyxl.load_workbook(args.input)
    ws = wb.active
    
    # The data columns are B through N (columns 2 through 14)
    # Headers are in row 1
    # We need to fill P2:P7 with formulas
    
    # Since TEXTJOIN might not be available, we'll build a manual concatenation
    # that handles commas properly
    
    for row_idx in range(2, 8):  # Rows 2-7
        # Build a complex formula that concatenates matching headers with commas
        # We'll use nested IF statements to build the result string
        
        # Strategy: For each column, if it matches, add it with a comma prefix
        # Then remove the leading comma at the end
        
        parts = []
        for col_idx in range(2, 15):  # Columns B through N
            col_letter = openpyxl.utils.get_column_letter(col_idx)
            header_cell = f"${col_letter}$1"
            value_cell = f"{col_letter}{row_idx}"
            min_cell = f"$O{row_idx}"
            
            # IF(AND(value=min, value>0), ","&header, "")
            parts.append(f'IF(AND({value_cell}={min_cell},{value_cell}>0),","&{header_cell},"")')
        
        # Concatenate all parts and remove the leading comma
        concat_formula = "&".join(parts)
        
        # Use MID to remove the first character (comma) if the result is not empty
        # MID(text, start, length) - start at position 2 to skip first comma
        # But we need to handle the case where result might be empty
        full_concat = f"({concat_formula})"
        
        # If the concatenated result is not empty, remove the first character (comma)
        # Otherwise return empty string
        formula = f'=IF({full_concat}<>"",MID({full_concat},2,LEN({full_concat})),"")'
        
        ws[f'P{row_idx}'] = formula
    
    # Save the workbook
    wb.save(args.output)
    print(f"Solution saved to {args.output}")


if __name__ == '__main__':
    main()
